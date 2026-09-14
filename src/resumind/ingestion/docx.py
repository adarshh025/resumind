import logging
from pathlib import Path
import docx
from docx.opc.exceptions import PackageNotFoundError

from resumind.ingestion.base import DocumentExtractor
from resumind.ingestion.models import DocumentExtractionResult

logger = logging.getLogger(__name__)

class DocxExtractor(DocumentExtractor):
    def extract(self, filepath: Path) -> DocumentExtractionResult:
        result = DocumentExtractionResult(
            filename=filepath.name,
            source_format="docx",
            raw_text="",
            status="SUCCESS"
        )
        
        try:
            doc = docx.Document(filepath)
            extracted_blocks = []
            
            # Extract paragraphs
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    extracted_blocks.append(text)
                    
            # Extract tables
            for table in doc.tables:
                extracted_blocks.append("--- TABLE_START ---")
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        # Extract cell text, handle line breaks inside cells
                        cell_text = cell.text.strip().replace('\n', ' ')
                        if cell_text and cell_text not in row_text: # simple dedup in row
                            row_text.append(cell_text)
                    if row_text:
                        extracted_blocks.append(" | ".join(row_text))
                extracted_blocks.append("--- TABLE_END ---")
                
            if not extracted_blocks:
                result.status = "TEXT_NOT_EXTRACTABLE"
                result.warnings.append("DOCX contains no extractable text.")
            else:
                result.raw_text = "\n".join(extracted_blocks)
                
        except PackageNotFoundError:
            result.status = "ERROR"
            result.errors.append("Invalid DOCX file or corrupted ZIP container.")
            logger.error(f"Failed to extract {filepath.name}: PackageNotFoundError")
        except Exception as e:
            result.status = "ERROR"
            result.errors.append(f"Unexpected extraction error: {str(e)}")
            logger.error(f"Failed to extract {filepath.name}: Unexpected error")
            
        return result
